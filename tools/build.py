"""Generate derived files from shared/: backend Schema.gs + Logic.gs, docs logic/sample data, Excel templates.

Run from anywhere:  python3 tools/build.py
"""
import json
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SHARED = ROOT / "shared"
schema = json.loads((SHARED / "schema.json").read_text())

# Sample data is clearly marked SAMPLE so it is never mistaken for real heater specs.
MODELS = [
    {"ModelID": "HM-01", "ModelName": "SAMPLE Heater 650-1000", "Make": "SAMPLE", "BeltWidthMin_mm": 650, "BeltWidthMax_mm": 1000, "RatingMax_Nmm": 1000, "PliesMax": 4, "Power_kW": "", "Voltage_V": 415, "Notes": "Replace with real data"},
    {"ModelID": "HM-02", "ModelName": "SAMPLE Heater 1000-1400", "Make": "SAMPLE", "BeltWidthMin_mm": 1000, "BeltWidthMax_mm": 1400, "RatingMax_Nmm": 1600, "PliesMax": 5, "Power_kW": "", "Voltage_V": 415, "Notes": "Replace with real data"},
    {"ModelID": "HM-03", "ModelName": "SAMPLE Heater 1400-2000", "Make": "SAMPLE", "BeltWidthMin_mm": 1400, "BeltWidthMax_mm": 2000, "RatingMax_Nmm": 2500, "PliesMax": 6, "Power_kW": "", "Voltage_V": 415, "Notes": "Replace with real data"},
]
PARTS = [("Top platen", 1, "Y"), ("Bottom platen", 1, "Y"), ("Control panel", 1, "Y"), ("Pressure bag", 1, "Y"), ("Power cable", 2, "Y"), ("Thermocouple", 2, "Y"), ("Side bar", 2, "Y"), ("Clamp bolt set", 1, "N")]
SETS_PER_MODEL = {"HM-01": 3, "HM-02": 4, "HM-03": 2}
TYPES = [{"ModelID": m["ModelID"], "ComponentName": n, "QtyPerSet": q, "IsSerialised": s, "UnitCost": ""} for m in MODELS for n, q, s in PARTS]

CONVEYORS = [
    {"ConveyorName": "SAMPLE CNV-101", "Area": "Raw Material", "BeltWidth_mm": 900, "BeltRating_Nmm": 800, "Plies": 4, "Notes": ""},
    {"ConveyorName": "SAMPLE CNV-205", "Area": "Sinter", "BeltWidth_mm": 1200, "BeltRating_Nmm": 1250, "Plies": 4, "Notes": ""},
    {"ConveyorName": "SAMPLE CNV-310", "Area": "Blast Furnace", "BeltWidth_mm": 1600, "BeltRating_Nmm": 2000, "Plies": 5, "Notes": ""},
]


def build_components():
    rows = []
    for m in MODELS:
        for set_no in range(1, SETS_PER_MODEL[m["ModelID"]] + 1):
            for name, qty, serial in PARTS:
                code = "".join(w[0] for w in name.split()).upper()
                units = range(1, qty + 1) if serial == "Y" else [None]
                for u in units:
                    cid = f"{m['ModelID']}-S{set_no}-{code}" + (f"{u}" if qty > 1 else "")
                    rows.append({"ComponentID": cid, "ModelID": m["ModelID"], "SetNo": set_no, "ComponentName": name, "SerialNo": "",
                                 "Qty": 1 if serial == "Y" else qty, "Condition": "OK", "Status": "In store", "MaintenanceStatus": "OK",
                                 "JointsSinceMaintenance": 0, "CurrentTicket": "", "CurrentConveyor": "", "LastInspected": "", "Remarks": ""})
    # Demo state: one HM-01 pressure bag damaged.
    for r in rows:
        if r["ComponentID"] == "HM-01-S3-PB":
            r.update(Condition="Damaged", Remarks="SAMPLE: torn bag")
    return rows


CONSUMABLES = ["SAR Jointing Kit", "SHR Jointing Kit", "UHR Jointing Kit", "Rema Tip Top SC2000", "Rema Tip Top Filler rubber",
               "Rubber strip 440 mm", "Rubber strip 220 mm", "Devcon putty", "R-flex putty", "T2 compound", "TN 9000 Solution",
               "TC 310 Cold solution", "TBS 3001 Solution", "WX101 Cold solution"]
CONS_ROWS = [{"ItemID": f"C{i:02d}", "ItemName": n, "Manufacturer": "", "Unit": "", "ReorderLevel": "", "SAPStock": "", "SAPStockUpdatedOn": ""}
             for i, n in enumerate(CONSUMABLES, 1)]

# Demo stock and six months of issues so the consumables screens have something to show. All marked SAMPLE.
CONS_BATCHES = [
    {"ItemID": "C01", "BatchNo": "SAMPLE-A1", "Qty": 6, "ExpiryDate": "2027-03-31", "ReceivedOn": "2026-04-02"},
    {"ItemID": "C01", "BatchNo": "SAMPLE-A2", "Qty": 2, "ExpiryDate": "2026-10-25", "ReceivedOn": "2026-01-15"},
    {"ItemID": "C04", "BatchNo": "SAMPLE-B1", "Qty": 14, "ExpiryDate": "2027-06-30", "ReceivedOn": "2026-07-01"},
    {"ItemID": "C08", "BatchNo": "SAMPLE-D1", "Qty": 3, "ExpiryDate": "2026-09-30", "ReceivedOn": "2025-12-10"},
    {"ItemID": "C11", "BatchNo": "SAMPLE-E1", "Qty": 20, "ExpiryDate": "2027-01-31", "ReceivedOn": "2026-08-05"},
]
CONS_TXNS = [
    {"Date": f"2026-{m:02d}-{d:02d}", "ItemID": item, "Type": "Issue", "Qty": q, "BatchNo": "SAMPLE", "Department": dept, "Conveyor": "", "Ref": "SAMPLE", "EnteredBy": "admin"}
    for m, d, item, q, dept in [
        (5, 6, "C01", 2, "FM"), (5, 20, "C04", 3, "FM"), (6, 3, "C01", 1, "FM"), (6, 18, "C11", 4, "FM"),
        (7, 9, "C01", 3, "FM"), (7, 22, "C04", 2, "FM"), (8, 5, "C01", 2, "FM"), (8, 27, "C11", 5, "FM"),
        (9, 4, "C01", 1, "FM"), (9, 16, "C04", 4, "FM"), (9, 29, "C08", 1, "FM"), (10, 2, "C01", 2, "FM")]
]

SAMPLE = {"HeaterModels": MODELS, "ComponentTypes": TYPES, "Components": build_components(), "Conveyors": CONVEYORS,
          "Departments": [{"Department": d} for d in ["FM", "RTG"]],
          "Consumables": CONS_ROWS, "ConsumableBatches": CONS_BATCHES, "ConsumableTransactions": CONS_TXNS}

GUIDE = {
    "HeaterModels": "One row per heater TYPE. Width range and max rating decide which belts it suits.",
    "ComponentTypes": "What one complete set of each model contains. QtyPerSet = how many of that part per set. IsSerialised = Y if each piece has its own ID.",
    "Components": "One row per PHYSICAL part in the store (or one row per lot for non-serialised items, with Qty). ComponentID is the sticker/QR number. Status: In store / Issued / Missing. Condition: OK / Damaged / Under repair / Scrapped.",
    "Conveyors": "Conveyor list for dropdowns. Belt details let the app suggest heaters automatically.",
    "Departments": "Departments that consume material.",
    "Consumables": "The 14 consumables. Fill Manufacturer, Unit (kit/kg/litre/roll), ReorderLevel and SAP stock.",
    "ConsumableBatches": "Current physical stock by batch: one row per batch with quantity and expiry.",
    "Users": "People who will log in. Role: Admin / Local / HOD / RTG Chief / FM Chief. Team: RTG / FM / Vendor. Leave password columns empty.",
}
TEMPLATE_TABS = ["HeaterModels", "ComponentTypes", "Components", "Conveyors", "Departments", "Consumables", "ConsumableBatches", "Users"]
HIDE_COLS = {"Users": ["UserID", "PasswordHash", "Salt", "MustChangePassword", "CreatedOn"]}


def write_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "RTG Store Heater & Consumables: data templates"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Fill each tab and send the file back. Rows marked SAMPLE are examples; overwrite or delete them. Keep the header row unchanged."
    r = 5
    for tab in TEMPLATE_TABS:
        ws.cell(r, 1, tab).font = Font(bold=True)
        ws.cell(r, 2, GUIDE[tab]).alignment = Alignment(wrap_text=True)
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    head_fill = PatternFill("solid", fgColor="1F3A68")
    for tab in TEMPLATE_TABS:
        cols = [c for c in schema[tab] if c not in HIDE_COLS.get(tab, [])]
        sh = wb.create_sheet(tab)
        for i, c in enumerate(cols, 1):
            cell = sh.cell(1, i, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            sh.column_dimensions[get_column_letter(i)].width = max(14, len(c) + 4)
        for ri, row in enumerate(SAMPLE.get(tab, []), 2):
            for ci, c in enumerate(cols, 1):
                sh.cell(ri, ci, row.get(c, ""))
        sh.freeze_panes = "A2"
    wb.save(path)


def main():
    schema_js = "// GENERATED by tools/build.py from shared/schema.json. Do not edit by hand.\nvar SCHEMA = " + json.dumps(schema, indent=2) + ";\n"
    (ROOT / "backend" / "Schema.gs").write_text(schema_js)
    sample_js = "// GENERATED by tools/build.py. Demo data only.\nvar SAMPLE_DATA = " + json.dumps(SAMPLE, indent=1) + ";\n"
    (ROOT / "backend" / "SampleData.gs").write_text(sample_js)
    logic = (SHARED / "logic.js").read_text()
    banner = "// GENERATED copy of shared/logic.js. Edit the shared file and run tools/build.py.\n"
    (ROOT / "backend" / "Logic.gs").write_text(banner + logic)
    shutil.copyfile(SHARED / "logic.js", ROOT / "docs" / "assets" / "js" / "logic.js")
    write_template(ROOT / "templates" / "RTG_data_templates.xlsx")
    # One file to paste into the Apps Script editor instead of nine.
    parts = sorted(p for p in (ROOT / "backend").glob("*.gs"))
    bundle = "// RTG Heater Store backend, all files combined. GENERATED by tools/build.py.\n\n" + "\n\n".join(
        f"// ===== {p.name} =====\n" + p.read_text() for p in parts)
    (ROOT / "dist").mkdir(exist_ok=True)
    (ROOT / "dist" / "Code.gs").write_text(bundle)
    # Demo mode runs this same backend in the browser.
    (ROOT / "docs" / "assets" / "js" / "backend-bundle.js").write_text(
        "// GENERATED by tools/build.py: the backend source, run in the browser for demo mode.\nwindow.RTG_BACKEND_SOURCE = " + json.dumps(bundle) + ";\n")
    print("built: dist/Code.gs, Schema.gs, SampleData.gs, Logic.gs, docs logic/sample data, templates/RTG_data_templates.xlsx")


if __name__ == "__main__":
    main()
